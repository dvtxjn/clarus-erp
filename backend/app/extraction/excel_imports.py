"""
Excel inputs to Module 3 — ported from reference/be_expense_sheet.py v7.2:
  - load_challan_due_amounts: daily duty challan list -> {BE no: Due Amount} (spec §5.1a)
  - load_org_repo: 'Organization List' export -> {AD Code: org name} (spec §5.1)
plus the two small calculations write_excel() did with them (name-mismatch
flag, interest split), pulled out so they don't depend on fixed invoice cells.
"""
from __future__ import annotations

import re
from datetime import datetime, timezone
from typing import BinaryIO, Mapping, Optional, Tuple, Union

from openpyxl import load_workbook

XlsxSource = Union[str, BinaryIO]


def challan_list_time(source: XlsxSource) -> Optional[datetime]:
    """When the challan list was generated: ICEGATE writes the .xlsx with its creation time (UTC).
    Re-saving in Excel keeps it. None when the file carries no time."""
    try:
        created = load_workbook(source, read_only=True).properties.created
    except Exception:
        return None
    finally:
        if hasattr(source, "seek"):
            source.seek(0)
    return created.replace(tzinfo=timezone.utc) if created and not created.tzinfo else created


def load_challan_rows(source: XlsxSource) -> Tuple[list[dict], Optional[str]]:
    """Daily duty challan list. Columns: IEC, Location Code, Doc type, Doc no.,
    Doc date, Challan no., Due Amount. Returns ([{be_no, be_date, location_code,
    iec, challan_no, due_amount}], note) — BE rows only; note is None on a clean
    read, else an error/warning for the user. A BE listed twice keeps its last row."""
    try:
        wb = load_workbook(source, data_only=True)
        ws = wb.active
    except Exception as e:
        return [], f"Couldn't read challan file: {e}"

    try:
        header_row = next(ws.iter_rows(min_row=1, max_row=1))
    except StopIteration:
        return [], "Challan file appears to be empty."
    header = [str(c.value or "").strip().lower() for c in header_row]
    col = {name: idx for idx, name in enumerate(header)}

    def pick(*names):
        return next((col[n] for n in names if n in col), None)

    doc_type_col = pick("doc type")
    doc_no_col = pick("doc no.", "doc no")
    due_col = pick("due amount")
    extra = {"be_date": pick("doc date"), "location_code": pick("location code"), "iec": pick("iec"),
             "challan_no": pick("challan no.", "challan no")}
    if doc_no_col is None or due_col is None:
        return [], "Couldn't find 'Doc no.' / 'Due Amount' columns in this file."

    def text(v) -> Optional[str]:
        if v is None:
            return None
        if isinstance(v, float) and v.is_integer():
            v = int(v)
        if hasattr(v, "strftime"):
            return v.strftime("%d.%m.%Y")
        return str(v).strip() or None

    rows: dict[str, dict] = {}
    dupes = set()
    for row in ws.iter_rows(min_row=2):
        if doc_type_col is not None:
            doc_type = str(row[doc_type_col].value or "").strip().upper()
            if doc_type and doc_type != "BE":
                continue  # e.g. shipping bills
        be_no = text(row[doc_no_col].value)
        try:
            due = float(row[due_col].value)
        except (TypeError, ValueError):
            continue
        if not be_no:
            continue
        if be_no in rows and rows[be_no]["due_amount"] != due:
            dupes.add(be_no)
        rows[be_no] = {"be_no": be_no, "due_amount": due,
                       **{k: (text(row[i].value) if i is not None else None) for k, i in extra.items()}}

    note = None
    if dupes:
        note = f"Note: {len(dupes)} BE number(s) appeared more than once with different amounts -- used the last one found."
    return list(rows.values()), note


def load_challan_due_amounts(source: XlsxSource) -> Tuple[dict[str, float], Optional[str]]:
    """{be_no: Due Amount} from the daily duty challan list (see load_challan_rows)."""
    rows, note = load_challan_rows(source)
    return {r["be_no"]: r["due_amount"] for r in rows}, note


def _cell(v) -> Optional[str]:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    t = str(v).strip()
    return t or None


def load_org_details(source: XlsxSource) -> list[dict]:
    """Every organisation in the filing software's 'Organization List' export
    (header row found in the first 5 rows), with all the details the invoice needs:
      name, short_names (ALIAS), ad_code (7 digits or None), gstin, pan, iec, cin,
      address (Branch AD1-3, City, State, Postal Code, Country), state, email, phone,
      is_active.
    Rows without an AD code are included too. Keyed by name (group companies can share
    an AD code); a name on several branch rows is taken once — the row with the most details."""
    wb = load_workbook(source, data_only=True, read_only=True)
    if "Organization List" not in wb.sheetnames:
        return []
    rows = list(wb["Organization List"].iter_rows(values_only=True))
    header_idx = next((r for r, vals in enumerate(rows[:5]) if "Organization" in vals), None)
    if header_idx is None:
        return []
    col = {str(v).strip().lower(): c for c, v in enumerate(rows[header_idx]) if v is not None}

    def get(vals, *names):
        for n in names:
            c = col.get(n.lower())
            if c is not None and c < len(vals):
                v = _cell(vals[c])
                if v:
                    return v
        return None

    out: dict[str, dict] = {}
    for vals in rows[header_idx + 1:]:
        name = get(vals, "Organization")
        if not name:
            continue
        ad = get(vals, "AD Code")
        raw_ad = vals[col["ad code"]] if "ad code" in col and col["ad code"] < len(vals) else None
        if ad and isinstance(raw_ad, (int, float)) and ad.isdigit():
            ad = ad.zfill(7)  # stored as a number, Excel dropped the leading zeros
        ad = ad if ad and re.fullmatch(r"\d{7}", ad) else None
        state = get(vals, "State")
        city, postal, country = get(vals, "City"), get(vals, "Postal Code"), get(vals, "Country")
        street = ", ".join(x.strip(" ,") for x in (get(vals, "Branch AD1"), get(vals, "Branch AD2"), get(vals, "Branch AD3")) if x)
        tail = ", ".join(x for x in (city, (state + (f" - {postal}" if postal else "")) if state else postal,
                                     country if country and country.upper() != "INDIA" else None) if x)
        org = {
            "name": name,
            "short_names": get(vals, "ALIAS"),
            "ad_code": ad,
            "gstin": (get(vals, "GSTIN") or "").upper() or None,
            "pan": (get(vals, "PAN NO", "PAN") or "").upper() or None,
            "iec": get(vals, "IE CODE NO", "IEC"),
            "address": ", ".join(x for x in (street, tail) if x) or None,
            "state": state,
            "email": get(vals, "Email Address", "Contact Email"),
            "phone": get(vals, "Telephone No", "Contact Mobile", "Contact Tel"),
            "is_active": (get(vals, "Is Active") or "True").upper() != "FALSE",
        }
        key = name.upper()
        filled = sum(1 for v in org.values() if v)
        if key not in out or filled > sum(1 for v in out[key].values() if v):
            out[key] = org
    return list(out.values())


def load_org_repo(source: XlsxSource) -> dict[str, str]:
    """AD Code -> Organization name (active rows with a 7-digit AD code) — the
    registry the BE name-mismatch check uses. Full details: load_org_details()."""
    return {o["ad_code"]: o["name"] for o in load_org_details(source)
            if o["ad_code"] and re.fullmatch(r"\d{7}", o["ad_code"]) and o["is_active"]}


def importer_name_check(ad_code: Optional[str], importer_name: Optional[str],
                        registry: Mapping[str, str]) -> Tuple[str, str]:
    """(mismatch_flag, billing_name). Same rule as write_excel():
    registered name contained in the BE name -> 'NO' (use BE name);
    otherwise 'YES' (bill under the registered name); AD code unknown -> 'N/A'."""
    if ad_code and ad_code in registry:
        registered = registry[ad_code]
        if registered.upper() in (importer_name or "").upper():
            return "NO", importer_name or registered
        return "YES", registered
    return "N/A", "AD Code Not Found"


def split_interest(be_duty_total: Optional[float], challan_due: float) -> Tuple[float, float]:
    """With interest applicable, the challan Due Amount (duty+GST+interest)
    replaces the BE's duty total as the payable figure.
    Returns (final_payable, interest_only)."""
    return challan_due, challan_due - (be_duty_total or 0)
