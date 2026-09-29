"""Excel copy of the proforma, styled like the client's template (Arial, navy
section bars, blue header cells, ₹ format) — but with dynamic charge rows and
plain numbers instead of the template's fixed-cell formulas."""
from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.invoice.company import BAR, BRAND, HEADER, HIGHLIGHT, SUBTOTAL

RUPEE = '"₹"#,##0.00;("₹"#,##0.00);"-"'
DARK = "1A1A1A"
GREY = "595959"
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
COLS = 6  # A..F, like the template


def _d(v):
    return float(Decimal(v)) if v not in (None, "") else None


def _fmt_date(iso):
    return date.fromisoformat(iso).strftime("%d-%b-%Y") if iso else ""


def _org_lines(org) -> list[str]:
    """Bill To details from the organization repository."""
    if not org:
        return []
    ids = "   ".join(f"{k}: {org[f]}" for k, f in (("GSTIN", "gstin"), ("PAN", "pan"), ("State", "state")) if org.get(f))
    contact = "   ".join(x for x in (org.get("email"), org.get("phone")) if x)
    return [x for x in (org.get("address"), ids, contact) if x]


VALUE_HEADS = ("Value of Goods", "GST Input", "Value / Kg", "Bill Rate (per kg)", "GST Output", "GST Difference")


def _value_block(ws, r: int, v: dict) -> int:
    """Template row 37-38: value of goods / GST input / value per kg / bill rate / GST output / difference."""
    from app.invoice.build import VALUE_NOTE, value_label

    vals = [v["value_of_goods"], v["gst_input"], v["value_per_kg"], v["bill_rate"], v["gst_output"], v["gst_difference"]]
    for col, (h, val) in enumerate(zip((value_label(v), *VALUE_HEADS[1:]), vals), start=1):
        hc = ws.cell(r, col, h)
        hc.font = Font(name="Arial", size=8.5, bold=True, color=DARK)
        hc.fill = PatternFill("solid", fgColor=HEADER)
        hc.alignment = Alignment(horizontal="center", wrap_text=True)
        hc.border = BOX
        vc = ws.cell(r + 1, col, _d(val))
        vc.font = Font(name="Arial", size=9.5, color=DARK)
        vc.alignment = Alignment(horizontal="right")
        vc.border = BOX
        vc.number_format = RUPEE if val is not None else "@"
    ws.merge_cells(start_row=r + 2, start_column=1, end_row=r + 2, end_column=COLS)
    note = ws.cell(r + 2, 1, VALUE_NOTE)
    note.font = Font(name="Arial", size=8, italic=True, color=GREY)
    return r + 4


_LOGO_PNG: bytes | None = None


def logo_png() -> bytes:
    """The Clarus logo as a PNG (drawn with the PDF's vector logo, rasterised once)."""
    global _LOGO_PNG
    if _LOGO_PNG is None:
        import pypdfium2
        from reportlab.pdfgen import canvas
        from app.invoice.pdf import ClarusLogo
        logo = ClarusLogo(40)
        buf = io.BytesIO()
        c = canvas.Canvas(buf, pagesize=(logo.width + 2, logo.height + 2))
        logo.drawOn(c, 1, 1)
        c.save()
        page = pypdfium2.PdfDocument(buf.getvalue())[0]
        img = page.render(scale=3).to_pil()
        out = io.BytesIO()
        img.save(out, format="PNG")
        _LOGO_PNG = out.getvalue()
    return _LOGO_PNG


def render_xlsx(inv: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Proforma Invoice"
    # wide enough that amounts never show as ####; long text wraps at spaces
    for col, width in zip("ABCDEF", (22, 36, 17, 16, 15, 17)):
        ws.column_dimensions[col].width = width
    ws.sheet_view.showGridLines = False
    # print: exactly one A4 page
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.4

    r = 1

    def merged(text, *, size=9, bold=False, color=DARK, fill=None, align="left", height=None, wrap=False):
        nonlocal r
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=COLS)
        c = ws.cell(r, 1, text)
        c.font = Font(name="Arial", size=size, bold=bold, color=color)
        c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
        if fill:
            for col in range(1, COLS + 1):
                ws.cell(r, col).fill = PatternFill("solid", fgColor=fill)
        if height:
            ws.row_dimensions[r].height = height
        r += 1

    def bar(text):
        merged(text, bold=True, color="FFFFFF", fill=BAR)

    co = inv["company"]
    # logo top-left, company details right-aligned (like the PDF)
    from openpyxl.drawing.image import Image as XlImage
    logo = XlImage(io.BytesIO(logo_png()))
    logo.height, logo.width = 46, int(46 * logo.width / logo.height)
    ws.add_image(logo, "A1")
    merged(co["name"], size=17, bold=True, color=BRAND, align="right", height=26)
    merged(co["address"], color=GREY, align="right", wrap=True, height=24)
    merged(co["tax_line"], color=GREY, align="right")
    merged(co["contact_line"], color=GREY, align="right")
    r += 1
    title = inv["title"] + (f" — {inv['copy_label'].upper()}" if inv.get("copy_label") else "")
    merged(title, size=16, bold=True, color="FFFFFF", fill=BAR, align="center", height=26)
    if inv.get("disclaimer"):
        merged(inv["disclaimer"].upper(), size=11, bold=True, color=DARK, fill=HIGHLIGHT, align="center", height=20)
    r += 1

    # BILL TO (A-C) | SHIPMENT DETAILS (D-F)
    bt, det = inv["bill_to"], inv["details"]
    left = [("BILL TO", None), (bt["name"] or "", "name")] + [(x, None) for x in _org_lines(bt.get("organization"))] + [
            ("BL Consignee: " + (bt["bl_consignee"] or "—"), None),
            ("BE Importer: " + (bt["be_importer"] or "—"), None)]
    if bt.get("hss"):
        left.append((f"HSS: {bt['hss']['seller']} (seller) → {bt['hss']['buyer']} (buyer)", None))
    right = [("SHIPMENT DETAILS", None), ("Invoice Date", _fmt_date(det["invoice_date"])), ("BE No.", det["be_no"] or "—"),
             ("BE Date", _fmt_date(det["be_date"]) or "—"), ("Port", det["port"] or "—")]
    for i in range(max(len(left), len(right))):
        if i < len(left):
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
            text, kind = left[i]
            c = ws.cell(r, 1, text)
            if i == 0:
                c.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
                for col in (1, 2, 3):
                    ws.cell(r, col).fill = PatternFill("solid", fgColor=BAR)
            else:
                c.font = Font(name="Arial", size=11 if kind == "name" else 9, bold=kind == "name", color=DARK)
        if i < len(right):
            label, value = right[i]
            if i == 0:
                ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=6)
                ws.cell(r, 4, label).font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
                for col in (4, 5, 6):
                    ws.cell(r, col).fill = PatternFill("solid", fgColor=BAR)
            else:
                ws.cell(r, 4, label).font = Font(name="Arial", size=9.5, bold=True, color=DARK)
                ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=6)
                ws.cell(r, 5, value).font = Font(name="Arial", size=9.5, color=DARK)
        r += 1
    r += 1

    # SHIPMENT REFERENCE
    bar("SHIPMENT REFERENCE")
    ref = inv["reference"]
    heads = ["Assessable Value", "MBL", "HBL / HSS", "# Containers", "WT (KGS)", "Exam Applicable"]
    vals = [_d(ref["assessable_value"]), ref["mbl"], ref["hbl"] or ref["hss"], ref["containers"],
            _d(ref["weight_kgs"]), ref["exam_applicable"]]
    for col, (h, v) in enumerate(zip(heads, vals), start=1):
        hc = ws.cell(r, col, h)
        hc.font = Font(name="Arial", size=8.5, bold=True, color=DARK)
        hc.fill = PatternFill("solid", fgColor=HEADER)
        hc.alignment = Alignment(horizontal="center", wrap_text=True)
        hc.border = BOX
        vc = ws.cell(r + 1, col, v)
        vc.font = Font(name="Arial", size=9.5, color=DARK)
        vc.alignment = Alignment(horizontal="center", wrap_text=True)
        vc.border = BOX
        if col in (1, 5) and v is not None:
            vc.number_format = "#,##0"
    r += 3
    # value of goods / GST input / bill rate next to the assessable value (client, 2026-09-30)
    r = _value_block(ws, r, inv["value"])

    # CHARGES
    bar("CHARGES")
    grand_row = None
    last_counted = [x for x in inv["sections"] if x["counts_in_total"] and (x["lines"] or x["category"] != "royalty")][-1]
    for sec in inv["sections"]:
        if not sec["lines"] and (not sec["counts_in_total"] or sec["category"] == "royalty"):
            continue  # empty Cost Inclusion / Royalty: not printed
        # (no Rate x Qty column — the description spans B:C)
        for col, h in enumerate(["Category", "Description", None, "Amount", "GST", "Total"], start=1):
            c = ws.cell(r, col, h)
            c.font = Font(name="Arial", size=9.5, bold=True, color=DARK)
            c.fill = PatternFill("solid", fgColor=HEADER)
            c.alignment = Alignment(horizontal="center")
            c.border = BOX
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        r += 1
        for li in sec["lines"]:
            row = [sec["title"], li["description"], None, _d(li["amount"]), _d(li["gst_amount"]), _d(li["total"])]
            for col, v in enumerate(row, start=1):
                c = ws.cell(r, col, v)
                c.font = Font(name="Arial", size=9, color=DARK)
                c.border = BOX
                c.alignment = Alignment(horizontal="left" if col <= 3 else "right", vertical="center",
                                        wrap_text=col <= 3)
                if col >= 4:
                    c.number_format = RUPEE
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
            r += 1
        if not sec["lines"]:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=COLS)
            ws.cell(r, 1, f"{sec['title']}: no charges").font = Font(name="Arial", size=9, italic=True, color=GREY)
            r += 1
        label = f"Subtotal - {sec['title']}" + ("" if sec["counts_in_total"] else " (not included in the total)")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        c = ws.cell(r, 1, label)
        c.font = Font(name="Arial", size=9.5, bold=True, color=DARK)
        c.alignment = Alignment(horizontal="right")
        t = ws.cell(r, 6, _d(sec["subtotal"]))
        t.font = Font(name="Arial", size=9.5, bold=True, color=DARK)
        t.number_format = RUPEE
        for col in range(1, COLS + 1):
            ws.cell(r, col).fill = PatternFill("solid", fgColor=SUBTOTAL)
        r += 2
        if sec is last_counted:  # grand total after the last section that counts
            if Decimal(inv.get("round_off") or 0):  # grand total is rounded to the rupee
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
                lab = ws.cell(r, 1, "Round off")
                lab.font, lab.alignment = Font(name="Arial", size=9, color=GREY), Alignment(horizontal="right")
                ro = ws.cell(r, 6, _d(inv["round_off"]))
                ro.font, ro.number_format = Font(name="Arial", size=9, color=GREY), RUPEE
                r += 1
            grand_row = r
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
            g = ws.cell(r, 1, "GRAND TOTAL — " + inv["grand_total_label"].upper())
            g.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
            g.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
            v = ws.cell(r, 6, _d(inv["grand_total"]))
            v.font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
            v.number_format = RUPEE
            for col in range(1, COLS + 1):
                ws.cell(r, col).fill = PatternFill("solid", fgColor=BRAND)
            ws.row_dimensions[r].height = 22
            r += 2
    assert grand_row is not None

    bar("NOTES")
    for i, note in enumerate(inv["notes"]):
        merged(note, fill=HIGHLIGHT if i == 0 else None, wrap=True, height=None if i == 0 else 26)
    r += 1
    bar("BANK DETAILS FOR PAYMENT")
    for label, value in inv["bank"]:
        ws.cell(r, 1, label).font = Font(name="Arial", size=9, bold=True, color=DARK)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        ws.cell(r, 2, value).font = Font(name="Arial", size=9, color=DARK)
        r += 1
    ws.cell(r - len(inv["bank"]), 5, f"For {inv['company']['name']}").font = Font(name="Arial", size=9, bold=True)
    ws.cell(r - 1, 5, "Authorised Signatory").font = Font(name="Arial", size=9, color=GREY)
    ws.print_area = f"A1:{get_column_letter(COLS)}{r}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
