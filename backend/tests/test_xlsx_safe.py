import io

from openpyxl import Workbook, load_workbook

from app.core.xlsx_safe import defuse


def test_text_starting_with_equals_is_not_a_formula():
    wb = Workbook()
    wb.active.append(['=HYPERLINK("http://x","click")', "ACME", -5, "-not a formula"])
    defuse(wb)
    buf = io.BytesIO()
    wb.save(buf)
    ws = load_workbook(io.BytesIO(buf.getvalue())).active
    assert ws["A1"].data_type == "s" and ws["A1"].value.startswith("'=")
    assert ws["B1"].value == "ACME" and ws["C1"].value == -5 and ws["D1"].value == "-not a formula"
