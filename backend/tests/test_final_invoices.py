"""Final invoices from a proforma — modelled on the client's CL/200/26-27 (tax) and
RI/CL/200/26-27 (reimbursement) for BAHUBALI RUBBER (Rajasthan -> IGST)."""
from tests.conftest import be_pdf
from tests.test_invoice import _setup


def test_final_invoices_from_proforma(client, admin_headers):
    h = admin_headers
    client.post("/organizations", json={"name": "BAHUBALI RUBBER PRIVATE LIMITED", "gstin": "08AACCB2783N1ZA",
                                        "pan": "AACCB2783N", "address": "E 700, RIICO INDUSTRIAL AREA CHOPANKI, Bhiwadi - 301019"},
                headers=h)
    sid = client.post("/shipments", json={"mbl": "274483845/QDDR2607499", "job": "165", "consignee": "BAHUBALI RUBBER",
                                          "container": "1"}, headers=h).json()["id"]
    client.patch(f"/shipments/{sid}", json={"duty_amount": "1436691", "igst_amount": "0", "port": "INGHR6"}, headers=h)
    pid = client.post(f"/shipments/{sid}/proformas", headers=h).json()["id"]
    charges = {c["code"]: c["id"] for c in client.get("/charge-master", headers=h).json()}
    client.post(f"/proformas/{pid}/line-items", json={"charge_master_id": charges["OTHERCHARGES"], "rate": 9500, "quantity": 1}, headers=h)

    r = client.post(f"/proformas/{pid}/final-invoices", headers=h)
    assert r.status_code == 201
    inv = {i["kind"]: i for i in r.json()}
    tax, reim = inv["tax"], inv["reimbursement"]
    # tax invoice: only Billed by Clarus — Agency 7,000 x 1 + Other 9,500; IGST (Rajasthan)
    assert [l["description"] for l in tax["lines"]] == ["Agency Charges", "Other Charges"]
    assert tax["intra_state"] is False and tax["place_of_supply"] == "[8] Rajasthan"
    assert tax["totals"]["before_tax"] == "16500.00" and tax["totals"]["gst"] == "2970.00"
    assert tax["totals"]["net_payable"] == "19470.00"
    assert tax["totals"]["in_words"] == "Nineteen Thousand Four Hundred Seventy Only."
    assert tax["header"]["job_number"].startswith("IMP/0165/") and tax["header"]["mbl_no"] == "274483845"
    assert tax["header"]["hbl_no"] == "QDDR2607499"
    assert [k for k, _ in tax["header_fields"]] == ["be_no", "be_date", "mbl_no", "hbl_no", "no_of_containers", "origin_port"]
    # reimbursement: customs duty as pure agent, no GST
    assert [(l["description"], l["tax_type"]) for l in reim["lines"]][0] == ("Customs Duty", "P")
    assert reim["totals"]["gst"] == "0.00" and reim["totals"]["net_payable"] == "1436691.00"
    assert reim["totals"]["in_words"] == "Fourteen Lakh Thirty Six Thousand Six Hundred Ninety One Only."

    # manual override, then issue: the pair shares the number
    t = client.patch(f"/final-invoices/{tax['id']}", json={"header": {"origin_port": "Qingdao"},
                     "lines": [*tax["lines"][:1], {**tax["lines"][1], "sub_description": "Transportation"}]}, headers=h).json()
    assert t["header"]["origin_port"] == "Qingdao" and t["lines"][1]["sub_description"] == "Transportation"
    t = client.post(f"/final-invoices/{tax['id']}/issue", headers=h).json()
    # Customs Duty must be the OOC copy's total: no OOC attached yet -> can't issue (client, 2026-10-08)
    r2 = client.post(f"/final-invoices/{reim['id']}/issue", headers=h)
    assert r2.status_code == 400 and "OOC" in r2.json()["detail"]
    client.post(f"/shipments/{sid}/documents", data={"document_type": "ooc_bill_of_entry"},
                files={"file": ("ooc.pdf", be_pdf(be_no="7770165"), "application/pdf")}, headers=h)
    fed = client.patch(f"/final-invoices/{reim['id']}", json={"lines": reim["lines"]}, headers=h).json()
    assert fed["lines"][0]["non_gst_value"] == "61500.00"  # the OOC total replaces the tracker's figure
    r2 = client.post(f"/final-invoices/{reim['id']}/issue", headers=h).json()
    n = t["number"].split("/")[1]
    assert t["number"].startswith("CL/") and r2["number"] == f"RI/CL/{n}/{t['number'].split('/')[2]}"
    # issued: still alterable until the 10th of next month (client, 2026-09-30); IRN any time
    if __import__("os").environ.get("TEST_DATABASE_URL"):  # altering needs Postgres (the lock trigger's switch)
        assert client.patch(f"/final-invoices/{tax['id']}", json={"remarks": "x"}, headers=h).status_code == 200
    assert client.patch(f"/final-invoices/{tax['id']}", json={"irn": "abc123"}, headers=h).json()["irn"] == "abc123"
    pdf = client.get(f"/final-invoices/{tax['id']}.pdf", headers=h)
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"
    # can't re-create from a proforma with issued invoices
    assert client.post(f"/proformas/{pid}/final-invoices", headers=h).status_code == 400


def test_intra_state_splits_cgst_sgst(client, admin_headers):
    h = admin_headers
    client.post("/organizations", json={"name": "Maha Intra Traders", "gstin": "27AAAAA0000A1Z5"}, headers=h)
    sid = client.post("/shipments", json={"mbl": "INTRA0001", "consignee": "Maha Intra Traders", "container": "2"},
                      headers=h).json()["id"]
    pid = client.post(f"/shipments/{sid}/proformas", headers=h).json()["id"]
    tax = next(i for i in client.post(f"/proformas/{pid}/final-invoices", headers=h).json() if i["kind"] == "tax")
    ln = tax["lines"][0]
    assert tax["intra_state"] and ln["cgst"] == ln["sgst"] == "1260.00" and ln["igst"] is None


def test_no_reimbursement_still_issued_as_not_applicable_and_pair_issued_together(client, admin_headers):
    """Client, 2026-09-30: tax and reimbursement always share one number. With no charges paid
    by us, the reimbursement invoice is still issued — no charge heads, "BILL CANCELLED — NOT
    APPLICABLE" — and both are issued in one action; one PDF with both, or two."""
    import pypdfium2

    h = admin_headers
    client.post("/organizations", json={"name": "Solo Tax Traders", "gstin": "24AAAAA1111A1Z5"}, headers=h)
    sid = client.post("/shipments", json={"mbl": "SOLOTAX0001", "consignee": "Solo Tax Traders", "container": "1"},
                      headers=h).json()["id"]
    pid = client.post(f"/shipments/{sid}/proformas", headers=h).json()["id"]  # agency only: nothing reimbursed
    made = {i["kind"]: i for i in client.post(f"/proformas/{pid}/final-invoices", headers=h).json()}
    assert set(made) == {"tax", "reimbursement"}
    assert made["reimbursement"]["not_applicable"] is True and made["reimbursement"]["lines"] == []
    assert made["tax"]["not_applicable"] is False

    both = client.post(f"/proformas/{pid}/final-invoices/issue", headers=h).json()
    nums = {i["kind"]: i["number"] for i in both}
    n, fy = nums["tax"].split("/")[1], nums["tax"].split("/")[2]
    assert nums["reimbursement"] == f"RI/CL/{n}/{fy}"                    # the chain doesn't break
    assert all(i["status"] == "issued" for i in both)
    assert client.post(f"/proformas/{pid}/final-invoices/issue", headers=h).status_code == 400  # already issued

    pdf = client.get(f"/proformas/{pid}/final-invoices.pdf", headers=h)
    assert pdf.status_code == 200 and len(pypdfium2.PdfDocument(pdf.content)) == 2   # tax + reimbursement
    ri = next(i for i in both if i["kind"] == "reimbursement")
    text = pypdfium2.PdfDocument(client.get(f"/final-invoices/{ri['id']}.pdf", headers=h).content)[0].get_textpage().get_text_range()
    assert "BILL CANCELLED" in text and "NOT APPLICABLE" in text


import os

import pytest


@pytest.mark.skipif(not os.environ.get("TEST_DATABASE_URL"), reason="altering issued invoices needs Postgres")
def test_altering_an_issued_invoice(client, admin_headers, monkeypatch):
    """Client, 2026-09-30: issued bills can be altered (mistakes happen) until the 10th of the next
    month; if e-invoicing applies (admin setting), only after confirming the e-invoice isn't filed."""
    from datetime import date as real_date

    import app.routers.final_invoices as fi
    from app.invoice.final import alter_until
    from app.models.final_invoice import FinalInvoice

    assert alter_until(FinalInvoice(invoice_date=real_date(2026, 9, 17))) == real_date(2026, 10, 10)
    assert alter_until(FinalInvoice(invoice_date=real_date(2026, 12, 3))) == real_date(2027, 1, 10)

    h = admin_headers
    client.post("/organizations", json={"name": "Alter Traders", "gstin": "24AAAAA2222A1Z5"}, headers=h)
    sid = client.post("/shipments", json={"mbl": "ALTER0001", "consignee": "Alter Traders", "container": "1"},
                      headers=h).json()["id"]
    pid = client.post(f"/shipments/{sid}/proformas", headers=h).json()["id"]
    client.post(f"/proformas/{pid}/final-invoices", headers=h)
    tax = next(i for i in client.post(f"/proformas/{pid}/final-invoices/issue", headers=h).json() if i["kind"] == "tax")
    assert tax["alter_until"]

    # e-invoicing off (default): alter straight away, number kept
    r = client.patch(f"/final-invoices/{tax['id']}", json={"remarks": "corrected"}, headers=h)
    assert r.status_code == 200 and r.json()["remarks"] == "corrected" and r.json()["number"] == tax["number"]

    # e-invoicing on: asked first; filed = blocked; not filed = altered
    assert client.put("/settings/e_invoicing", json={"value": True}, headers=h).json()["e_invoicing"] is True
    ask = client.patch(f"/final-invoices/{tax['id']}", json={"remarks": "again"}, headers=h)
    assert ask.status_code == 409 and ask.json()["detail"]["ask"] == "e_invoice_filed"
    assert client.patch(f"/final-invoices/{tax['id']}", json={"remarks": "again", "e_invoice_filed": True},
                        headers=h).status_code == 400
    ok = client.patch(f"/final-invoices/{tax['id']}", json={"remarks": "again", "e_invoice_filed": False}, headers=h)
    assert ok.status_code == 200 and ok.json()["remarks"] == "again"
    # an IRN entered = filed
    client.patch(f"/final-invoices/{tax['id']}", json={"irn": "IRN123"}, headers=h)
    assert client.patch(f"/final-invoices/{tax['id']}", json={"remarks": "x", "e_invoice_filed": False},
                        headers=h).status_code == 400
    client.put("/settings/e_invoicing", json={"value": False}, headers=h)

    # after the 10th of next month: no longer
    class Later(real_date):
        @classmethod
        def today(cls):
            return real_date(2099, 1, 1)

    monkeypatch.setattr(fi, "date", Later)
    late = client.patch(f"/final-invoices/{tax['id']}", json={"remarks": "late"}, headers=h)
    assert late.status_code == 400 and "credit note" in late.json()["detail"]


def test_settings_are_admin_only(client, admin_headers):
    assert client.put("/settings/e_invoicing", json={"value": "yes"}, headers=admin_headers).status_code == 422
    assert client.put("/settings/nope", json={"value": True}, headers=admin_headers).status_code == 404
    assert client.get("/settings").status_code == 401


def test_invoice_register_filters_and_exports(client, admin_headers):
    """Client, 2026-09-30: one Invoices page — every tax / reimbursement invoice, filtered by year,
    type, status, client, number; printed as one PDF or exported as a register, without opening shipments."""
    import io

    import pypdfium2
    from openpyxl import load_workbook

    h = admin_headers
    for name, gst in (("Register Alpha Ltd", "24AAAAA3333A1Z5"), ("Register Beta Ltd", "24AAAAA4444A1Z5")):
        client.post("/organizations", json={"name": name, "gstin": gst}, headers=h)
        sid = client.post("/shipments", json={"mbl": f"REG{gst[7:11]}", "consignee": name, "container": "1", "job": gst[7:10]},
                          headers=h).json()["id"]
        pid = client.post(f"/shipments/{sid}/proformas", headers=h).json()["id"]
        client.post(f"/proformas/{pid}/final-invoices", headers=h)
        client.post(f"/proformas/{pid}/final-invoices/issue", headers=h)

    reg = client.get("/final-invoices", params={"client": "register alpha"}, headers=h).json()
    assert {r["kind"] for r in reg["invoices"]} == {"tax", "reimbursement"} and all(r["customer"] == "Register Alpha Ltd" for r in reg["invoices"])
    assert reg["financial_years"]
    tax_only = client.get("/final-invoices", params={"client": "register", "kind": "tax"}, headers=h).json()["invoices"]
    assert len(tax_only) == 2 and all(r["kind"] == "tax" for r in tax_only)
    one = client.get("/final-invoices", params={"q": tax_only[0]["number"]}, headers=h).json()["invoices"]
    # a number finds the pair (they share it): CL/n and RI/CL/n
    assert sorted(r["number"] for r in one) == sorted([tax_only[0]["number"], "RI/" + tax_only[0]["number"]])

    ids = ",".join(str(r["id"]) for r in client.get("/final-invoices", params={"client": "register"}, headers=h).json()["invoices"])
    pdf = client.get("/final-invoices/export.pdf", params={"ids": ids}, headers=h)
    assert pdf.status_code == 200 and len(pypdfium2.PdfDocument(pdf.content)) == 4
    xl = client.get("/final-invoices/register.xlsx", params={"client": "register"}, headers=h)
    ws = load_workbook(io.BytesIO(xl.content)).active
    assert ws.max_row == 5 and ws["A1"].value == "Number"
    assert client.get("/final-invoices").status_code == 401


def test_proforma_register(client, admin_headers):
    import pypdfium2

    h = admin_headers
    sid = client.post("/shipments", json={"mbl": "PFREG0001", "consignee": "Proforma Register Co", "job": "881"},
                      headers=h).json()["id"]
    pid = client.post(f"/shipments/{sid}/proformas", headers=h).json()["id"]
    rows = client.get("/proformas", params={"q": "881"}, headers=h).json()
    assert [r["id"] for r in rows["proformas"]] == [pid] and rows["proformas"][0]["status"] == "draft"
    assert rows["proformas"][0]["grand_total"] and rows["financial_years"]
    assert client.get("/proformas", params={"q": "881", "status": "sent"}, headers=h).json()["proformas"] == []
    pdf = client.get("/proformas/export.pdf", params={"ids": str(pid)}, headers=h)
    assert pdf.status_code == 200 and len(pypdfium2.PdfDocument(pdf.content)) == 1


def test_company_bank_and_terms_come_from_settings(client, admin_headers):
    """Client, 2026-09-30 (P0): the admin changes what's printed on invoices on the Settings page."""
    h = admin_headers
    s = client.get("/settings", headers=h).json()
    assert s["company"]["gstin"] == "27AAVFC6734E1Z4" and s["bank"][0][0] == "Account Name"
    co = {**s["company"], "name": "CLARUS LOGISTICS LLP (TEST)", "phone": "+91 00000 00000"}
    assert client.put("/settings/company", json={"value": co}, headers=h).status_code == 200
    assert client.put("/settings/company", json={"value": {**co, "gstin": "bad"}}, headers=h).status_code == 422
    client.put("/settings/bank", json={"value": [["Account Name", "CLARUS TEST"], ["Bank Name", "Test Bank"]]}, headers=h)
    client.put("/settings/proforma_notes", json={"value": ["A note from settings"]}, headers=h)

    sid = client.post("/shipments", json={"mbl": "SETTINGS0001", "consignee": "Settings Co"}, headers=h).json()["id"]
    pid = client.post(f"/shipments/{sid}/proformas", headers=h).json()["id"]
    inv = client.get(f"/proformas/{pid}/invoice", headers=h).json()
    assert inv["company"]["name"] == "CLARUS LOGISTICS LLP (TEST)" and "+91 00000 00000" in inv["company"]["contact_line"]
    assert inv["bank"] == [["Account Name", "CLARUS TEST"], ["Bank Name", "Test Bank"]]
    assert inv["notes"] == ["A note from settings"]
    # back to the real values for the other tests
    for k in ("company", "bank", "proforma_notes"):
        client.put(f"/settings/{k}", json={"value": s[k]}, headers=h)


def test_number_formats_restart_and_deleting_trial_invoices(client, admin_headers):
    """Client, 2026-09-30: trial invoices are deleted before going live, the series restarts from a new number
    in the admin's own format — and no number is ever given twice."""
    from app.invoice.final import fy_of
    from datetime import date

    h, fy = admin_headers, fy_of(date.today())

    def pair(mbl):
        client.post("/organizations", json={"name": f"Series Co {mbl}", "gstin": "24AAAAA2222A1Z5"}, headers=h)
        sid = client.post("/shipments", json={"mbl": mbl, "consignee": f"Series Co {mbl}", "container": "1"}, headers=h).json()["id"]
        pid = client.post(f"/shipments/{sid}/proformas", headers=h).json()["id"]
        client.post(f"/proformas/{pid}/final-invoices", headers=h)
        return pid

    # bad formats refused
    assert client.put("/invoice-series", json={"tax": "CL/{fy}", "reimbursement": "RI/{n}"}, headers=h).status_code == 422
    assert client.put("/invoice-series", json={"tax": "X/{n}", "reimbursement": "X/{n}"}, headers=h).status_code == 422
    assert client.put("/invoice-series", json={"tax": "CL/{bad}", "reimbursement": "RI/{n}"}, headers=h).status_code == 422
    r = client.put("/invoice-series", json={"tax": "CLR/{n:04d}/{fy}", "reimbursement": "CLR-RI/{n:04d}/{fy}"}, headers=h)
    assert r.status_code == 200 and r.json()["example_tax"] == f"CLR/0201/{fy}"

    assert client.put("/invoice-counter", json={"fy": fy, "next_seq": 5000}, headers=h).status_code == 200
    trial = client.post(f"/proformas/{pair('SERIES0001')}/final-invoices/issue", headers=h).json()
    nums = {i["kind"]: i["number"] for i in trial}
    assert nums == {"tax": f"CLR/5000/{fy}", "reimbursement": f"CLR-RI/5000/{fy}"}

    # the admin deletes the trial pair (issued) — restorable, number kept out of use
    for i in trial:
        assert client.delete(f"/final-invoices/{i['id']}", headers=h).status_code == 204
    # restart at the same number: refused, it was given once already
    client.put("/invoice-counter", json={"fy": fy, "next_seq": 5000}, headers=h)
    r = client.post(f"/proformas/{pair('SERIES0002')}/final-invoices/issue", headers=h)
    assert r.status_code == 400 and "already used" in r.json()["detail"]
    client.put("/invoice-counter", json={"fy": fy, "next_seq": 5001}, headers=h)
    live = client.post(f"/proformas/{pair('SERIES0003')}/final-invoices/issue", headers=h).json()
    assert {i["number"] for i in live} == {f"CLR/5001/{fy}", f"CLR-RI/5001/{fy}"}
    client.put("/invoice-series", json={"tax": "CL/{n}/{fy}", "reimbursement": "RI/CL/{n}/{fy}"}, headers=h)  # back for other tests


def test_customs_duty_is_always_the_ooc_total(client, admin_headers):
    """Client, 2026-10-08: Customs Duty on the reimbursement invoice = the OOC BE's total duty, no
    exception — not editable on the proforma or the final invoice; Stamp Duty stays editable."""
    from tests.test_invoice import _setup

    h = admin_headers
    sid = _setup(client, h, port_consignee="Divine")
    pid = client.post(f"/shipments/{sid}/proformas", headers=h).json()["id"]
    p = client.post(f"/proformas/{pid}/fill-from-shipment", headers=h).json()["proforma"]
    cd = next(li for li in p["line_items"] if li["description"].startswith("Customs Duty"))
    assert float(cd["total"]) == 61500.0
    r = client.patch(f"/proformas/{pid}/line-items/{cd['id']}", json={"rate": "1"}, headers=h)
    assert r.status_code == 400 and "OOC" in r.json()["detail"]

    ri = next(i for i in client.post(f"/proformas/{pid}/final-invoices", headers=h).json() if i["kind"] == "reimbursement")
    duty = next(ln for ln in ri["lines"] if ln["code"] == "CD")
    assert duty["non_gst_value"] == "61500.00"
    # typing another figure: Customs Duty goes back to the OOC total, Stamp Duty keeps the edit
    lines = [{**ln, "non_gst_value": "1" if ln["code"] in ("CD", "SD") else ln["non_gst_value"]} for ln in ri["lines"]]
    lines = [{k: ln[k] for k in ("description", "sac", "tax_type", "non_gst_value", "taxable_value", "gst_rate", "code")}
             for ln in lines]
    out = client.patch(f"/final-invoices/{ri['id']}", json={"lines": lines}, headers=h).json()
    by = {ln["code"]: ln for ln in out["lines"]}
    assert by["CD"]["non_gst_value"] == "61500.00" and by["SD"]["non_gst_value"] == "1.00"
